"""Turns an uploaded handbook (Word, PDF, Markdown or plain text) into Sigap's article format.

The formatting itself is done by the Langflow flow "handbook_formatter" (an LLM), one excerpt at a time. This
module only extracts the text, splits it, and checks every draft against the source, because a model that
rewrites rules can change a number or drop a condition:

  - a number in a drafted article that is not in the source excerpt, or a number in the source that no article
    carries, is reported;
  - source sentences that no article seems to cover are listed as "possibly missing";
  - translated lines are always flagged for review.

The result is a draft for an admin to review and edit; it is never activated automatically.
"""
import io
import json
import re
import threading
import time
import uuid
from datetime import date
from pathlib import Path

from .services import validate_handbook

CHUNK_CHARS = 6000
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
FORMATTER_FLOW = "handbook_formatter"
_ARTICLE_START = re.compile(r"^\s*(#+\s|pasal\b|bab\b|article\b|chapter\b|\d+(\.\d+)*[.)]?\s+[A-Z])", re.I)
_NUMBER = re.compile(r"(?<![\w.])\d+(?:[.,]\d+)?(?![\w])")
_NOT_A_RULE_NUMBER = re.compile(r"(?:pasal|bab|ayat|article|chapter|halaman|page|hal\.)\s*$", re.I)
# Words that occur in almost every question; as tags they give unrelated questions free hits.
_GENERIC_TAGS = {"mahasiswa", "aturan", "kampus", "universitas", "pasal", "peraturan", "student", "students", "rule",
                 "rules", "university", "campus", "article", "dan", "atau", "yang", "the", "and", "or", "mata",
                 "kuliah", "matkul", "course", "courses", "untuk", "dari", "dengan", "setiap", "saya", "aku", "for"}


class FormatError(ValueError):
    pass


# ---------------------------------------------------------------- reading the upload

def extract_text(filename: str, data: bytes) -> str:
    """Plain text of an uploaded document. Headings from Word keep a '#' so article boundaries survive."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise FormatError("Berkas lebih dari 8 MB. Pecah dokumen atau unggah bab yang berubah saja.")
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:  # pypdf raises many types for damaged files
            raise FormatError(f"PDF tidak bisa dibaca: {e}") from None
        if len(text.strip()) < 200:
            raise FormatError("PDF ini hampir tidak berisi teks (kemungkinan hasil pindaian). "
                              "Unggah versi Word atau PDF yang teksnya bisa disalin.")
        return text
    if ext == ".docx":
        import docx
        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise FormatError(f"Dokumen Word tidak bisa dibaca: {e}") from None
        lines = []
        for para in d.paragraphs:
            t = para.text.strip()
            if t:
                lines.append(("# " if (para.style.name or "").lower().startswith(("heading", "judul")) else "") + t)
        for table in d.tables:  # tables (e.g. penalty grids) come after the text, one row per line
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if ext in (".md", ".txt", ""):
        return data.decode("utf-8", errors="replace")
    raise FormatError("Format belum didukung. Unggah .docx, .pdf (berisi teks), .md, atau .txt.")


_TOC_LINE = re.compile(r"(\.{4,}|…{2,}|\s{3,})\s*\d+\s*$")  # "Pasal 4.2 Keterlambatan ........ 12"


def _is_heading(line: str) -> bool:
    letters = [c for c in line if c.isalpha()]
    return line.lstrip().startswith("#") or (len(letters) >= 4 and sum(c.isupper() for c in letters) / len(letters) > 0.7)


def split_excerpts(text: str, limit: int = CHUNK_CHARS) -> list[str]:
    """Excerpts of at most ~limit characters, cut before an article or chapter heading where possible.
    Table-of-contents lines are dropped first: they repeat article titles with page numbers."""
    lines = [l.rstrip() for l in text.replace("\r", "").split("\n") if not _TOC_LINE.search(l.rstrip())]
    out, cur, size = [], [], 0
    for line in lines:
        starts_article = bool(_ARTICLE_START.match(line))
        if cur and (size + len(line) > limit or (starts_article and size > limit * 0.6)):
            out.append("\n".join(cur).strip())
            cur, size = [], 0
        cur.append(line)
        size += len(line) + 1
    if "\n".join(cur).strip():
        out.append("\n".join(cur).strip())
    return [e for e in out if e]


# ---------------------------------------------------------------- the agent's reply

def parse_reply(text: str) -> list[dict]:
    """The JSON array the formatter returns, tolerating a code fence or text around it."""
    start, end = text.find("["), text.rfind("]")
    if start < 0 or end < start:
        raise FormatError("Agen tidak mengembalikan daftar pasal.")
    try:
        items = json.loads(text[start:end + 1])
    except json.JSONDecodeError as e:
        raise FormatError(f"Balasan agen bukan JSON yang valid: {e}") from None
    return [i for i in items if isinstance(i, dict)]


def _one_line(s) -> str:
    s = re.sub(r"\s*\n\s*[-*•]\s*", "; ", str(s or ""))  # list items the model kept become clauses
    return re.sub(r"\s+", " ", s).strip()


def _clean_tags(tags) -> list[str]:
    """Single words only: the search counts each word of a tag as a hit, so a phrase like "kelas karyawan" would
    let an unrelated question ("beasiswa kelas karyawan") reach the two-hit threshold with one concept.
    Contact details and long numbers (phones, emails, sites) are dropped: students do not ask with them."""
    words = tags if isinstance(tags, list) else str(tags or "").split(",")
    out = []
    for t in words:
        t = str(t).lower().strip().strip(".")
        if (not t or " " in t or "@" in t or "/" in t or re.search(r"\.\w{2,}", t) or re.search(r"\d{4,}", t)
                or re.fullmatch(r"\d+%", t) or t in _GENERIC_TAGS or len(t) < 2 or t in out):
            continue
        out.append(re.sub(r"[^\w-]", "", t))
    return [t for t in out if t][:20]


def _numbers(text: str) -> set[str]:
    """Rule-bearing numbers: '75', '24', '2,5'. Article/chapter/page numbers are left out."""
    found = set()
    for m in _NUMBER.finditer(text):
        if _NOT_A_RULE_NUMBER.search(text[max(0, m.start() - 12):m.start()]):
            continue
        line_start = text.rfind("\n", 0, m.start()) + 1
        if text[line_start:m.start()].strip() in ("", "(", "#", "##") and re.match(r"[.)]", text[m.end():m.end() + 1] or "x"):
            continue  # list markers like "(1)" or "2."
        found.add(m.group().replace(",", "."))
    return found


def _is_sentence(text: str) -> bool:
    """A rule text, not an identifier or a placeholder: several words with spaces."""
    return len(text) >= 25 and len(text.split()) >= 5


def _words(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-zà-ÿ]{4,}", text.lower())}


def _passages(source: str, articles: list[dict]) -> dict[str, str]:
    """Each article's own passage in the source: from its heading ("Pasal 4.2", "4.2 Judul") to the next one.
    Articles whose number the source never states are missing from the result."""
    starts = []
    for a in articles:
        m = re.search(r"(?mi)^\s*#*\s*(?:pasal|article)?\s*" + re.escape(a["number"]) + r"(?![\d.]*\d)", source)
        if m:
            starts.append((m.start(), a["number"]))
    starts.sort()
    return {num: source[pos:(starts[i + 1][0] if i + 1 < len(starts) else len(source))]
            for i, (pos, num) in enumerate(starts)}


def _rule_numbers(text: str) -> set[str]:
    # Numbers that only appear in headings (titles, years on a cover) are not rules.
    return _numbers("\n".join(l for l in text.split("\n") if not _is_heading(l)))


def _context(text: str, n: str) -> str:
    i = text.find(n)
    return _one_line(text[max(0, i - 50):i + 50]) if i >= 0 else ""


def check_excerpt(source: str, articles: list[dict]) -> list[str]:
    """Differences between an excerpt and the articles drafted from it, for the reviewer. Numbers are compared
    article by article with that article's own passage, so a 75 changed to 70 is caught even when 70 appears
    elsewhere in the handbook."""
    warnings = []
    passages = _passages(source, articles)
    for a in articles:
        own = passages.get(a["number"])
        scope = own if own is not None else source
        nums, src = _numbers(a["id"]), _rule_numbers(scope)
        extra = sorted(nums - src - {a["number"]})
        if extra:
            warnings.append(f"Pasal {a['number']}: angka {', '.join(extra)} tidak ada di "
                            f"{'pasal ini di ' if own is not None else ''}dokumen asli. Periksa isinya.")
        if own is not None:
            for n in sorted(src - nums - {a["number"]}, key=float):
                warnings.append(f"Pasal {a['number']}: angka {n} di dokumen asli tidak ada di draf "
                                f"(sekitar: \"…{_context(own, n)}…\").")
        if len(a["tags"]) < 6:
            warnings.append(f"Pasal {a['number']}: hanya {len(a['tags'])} tag; tambah kata yang dipakai mahasiswa.")
    # Text outside any recognised article (unnumbered documents, or text before the first article).
    rest = source
    for passage in passages.values():
        rest = rest.replace(passage, "\n")
    drafted = set().union(*(_numbers(a["id"]) for a in articles)) if articles else set()
    for n in sorted(_rule_numbers(rest) - drafted, key=float):
        warnings.append(f"Angka {n} di dokumen asli tidak muncul di draf (sekitar: \"…{_context(rest, n)}…\").")
    covered = [_words(a["id"] + " " + a["en"]) for a in articles]  # the source may be in either language
    for sentence in re.split(r"(?<=[.;:])\s+|\n+", source):
        sentence = sentence.strip()
        w = _words(sentence)
        if len(sentence) < 80 or len(w) < 6:
            continue
        best = max((len(w & c) / len(w) for c in covered), default=0)
        if best < 0.35:
            warnings.append(f"Mungkin terlewat: \"{_one_line(sentence)[:160]}\"")
    return warnings


# ---------------------------------------------------------------- the whole document

def format_document(text: str, meta: dict, run_flow, progress=lambda done, total: None) -> dict:
    """run_flow(input_text) -> reply text. Returns {"markdown", "articles", "warnings", "valid", "error"}."""
    excerpts = split_excerpts(text)
    if not excerpts:
        raise FormatError("Dokumen kosong.")
    articles, warnings, prev = [], [], "0"
    progress(0, len(excerpts))
    for i, excerpt in enumerate(excerpts, 1):
        reply = run_flow(f"Document: {meta.get('title_id', '')}\nPrevious article number: {prev}\n\nExcerpt:\n{excerpt}")
        try:
            raw = parse_reply(reply)
        except FormatError as e:
            warnings.append(f"Bagian {i} dari {len(excerpts)} gagal diformat ({e}); salin pasalnya secara manual.")
            raw = []
        chunk_articles = []
        for item in raw:
            number = str(item.get("number") or "").strip().removeprefix("Pasal ").strip()
            if not re.fullmatch(r"\d+(\.\d+)*", number):
                warnings.append(f"Pasal dengan nomor tidak valid ({number or 'kosong'}) di bagian {i}; diberi nomor baru.")
                number = f"{prev.split('.')[0] if '.' in prev else prev}.{len(chunk_articles) + 1}"
            a = {"number": number, "title_id": _one_line(item.get("title_id")).replace("|", "/"),
                 "title_en": _one_line(item.get("title_en")).replace("|", "/") or _one_line(item.get("title_id")),
                 "tags": _clean_tags(item.get("tags")), "id": _one_line(item.get("text_id")),
                 "en": _one_line(item.get("text_en"))}
            if not _is_sentence(a["id"]):
                warnings.append(f"Pasal {number}: agen tidak menulis isi aturan dalam Bahasa Indonesia "
                                f"(\"{a['id'][:60]}\"); salin isinya dari dokumen asli.")
                a["id"] = a["id"] or "[ISI PASAL]"
            if not _is_sentence(a["en"]):
                warnings.append(f"Pasal {number}: baris EN kosong atau bukan kalimat; terjemahkan secara manual.")
                a["en"] = a["en"] or "[TRANSLATION]"
            chunk_articles.append(a)
        warnings += check_excerpt(excerpt, chunk_articles)
        articles += chunk_articles
        if chunk_articles:
            prev = chunk_articles[-1]["number"]
        progress(i, len(excerpts))
    seen = set()
    for a in articles:
        if a["number"] in seen:
            warnings.append(f"Nomor pasal {a['number']} dipakai lebih dari sekali; beri nomor berbeda.")
        seen.add(a["number"])
    if not articles:
        raise FormatError("Agen tidak menemukan aturan di dokumen ini.")
    for i, a in enumerate(articles):  # articles that share many tags compete for the same questions
        for b in articles[i + 1:]:
            shared = sorted(set(a["tags"]) & set(b["tags"]))
            if len(shared) >= 4:
                warnings.append(f"Pasal {a['number']} dan {b['number']} berbagi tag ({', '.join(shared)}); pertanyaan bisa "
                                f"tertukar. Sisakan tag yang khas di masing-masing pasal.")
    warnings.insert(0, "Teks dan terjemahan dibuat oleh agen AI. Bandingkan dengan dokumen resmi sebelum mengaktifkan, "
                       "terutama angka, batas waktu, dan baris EN.")
    markdown = assemble(meta, articles)
    try:
        validate_handbook(markdown)
        valid, error = True, None
    except ValueError as e:
        valid, error = False, str(e)
    return {"markdown": markdown, "articles": len(articles), "warnings": warnings, "valid": valid, "error": error}


def assemble(meta: dict, articles: list[dict]) -> str:
    front = [f"{k}: {_one_line(meta[k])}" for k in ("title_id", "title_en", "institution", "version", "effective")
             if _one_line(meta.get(k))]
    front.append(f"drafted_by: agen format Sigap {date.today().isoformat()}, perlu ditinjau")
    body = [f"## {a['number']} | {a['title_id']} | {a['title_en']}\ntags: {', '.join(a['tags'])}\nID: {a['id']}\nEN: {a['en']}"
            for a in articles]
    return "---\n" + "\n".join(front) + "\n---\n\n" + "\n\n".join(body) + "\n"


# ---------------------------------------------------------------- background jobs (admin panel polls these)

_jobs: dict[str, dict] = {}
_lock = threading.Lock()


def start_job(text: str, meta: dict, run_flow) -> str:
    job_id = uuid.uuid4().hex[:12]
    with _lock:
        for k in [k for k, j in _jobs.items() if time.time() - j["created"] > 3600]:
            del _jobs[k]  # drafts are kept for an hour; the admin copies the result into the upload box
        _jobs[job_id] = {"status": "running", "done": 0, "total": 0, "created": time.time()}

    def progress(done, total):
        with _lock:
            _jobs[job_id].update(done=done, total=total)

    def work():
        try:
            result = format_document(text, meta, run_flow, progress)
            update = {"status": "done", "result": result}
        except Exception as e:  # shown to the admin instead of a silent failure
            update = {"status": "error", "error": str(e)}
        with _lock:
            _jobs[job_id].update(update)

    threading.Thread(target=work, daemon=True).start()
    return job_id


def job(job_id: str) -> dict | None:
    with _lock:
        j = _jobs.get(job_id)
        return {k: v for k, v in j.items() if k != "created"} if j else None
